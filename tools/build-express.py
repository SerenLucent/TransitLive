"""Extract stop/pass facts from official Korail paired arrival/departure sheets."""
import datetime, json, re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
lines = json.loads((ROOT/'lines.js').read_text(encoding='utf-8').split('=',1)[1].strip().removesuffix(';'))
aliases = {
 '디엠시':'디지털미디어시티','항공대':'한국항공대','효창공':'효창공원앞','홍대입':'홍대입구',
 '평내호':'평내호평','신인천':'인천','온양온':'온양온천','쌍용나':'쌍용','성균관':'성균관대',
 '금천구':'금천구청','서울역':'서울','지청':'청량리','동대문역':'동대문','지하서':'서울',
 '신길온':'신길온천','정부과':'정부과천청사','경마공':'경마공원','별내별':'별내별가람',
 '당고개':'불암산','총신대':'총신대입구','서울숲':'서울숲',
 '가산디':'가산디지털단지','동대':'동대문','동두중':'동두천중앙','종로':'종로3가','종로5':'종로5가','평지제':'평택지제',
 '과천청':'정부과천청사','동운':'동대문역사문화공원','미아4':'미아사거리','별가람':'별내별가람',
 '성신여':'성신여대입구','숙대입':'숙대입구','충무':'충무로','평촌동':'평촌','한성대':'한성대입구',
 '강남구':'강남구청','구룡역':'구룡','남동인':'남동인더스파크','대모산':'대모산입구',
 '로데오':'압구정로데오','매탄권':'매탄권선','소래포':'소래포구','수원시':'수원시청','신수원':'수원','인천논':'인천논현',
}
def norm(value):
 value=re.sub(r'\s|\([^)]*\)','',str(value or ''))
 value=re.sub(r'^\d+','',value)
 return aliases.get(value,value)
def is_time(v):
 return isinstance(v,(datetime.time,datetime.datetime)) or bool(re.fullmatch(r'\d{1,2}:\d{2}(:\d{2})?',str(v or '').strip()))

sources=[
 ('korail-weekday-20260901.xlsx','weekday',None),('korail-holiday-20260901.xlsx','holiday',None),
 ('korail-line1-weekday-20260901.xlsx','weekday',1),('korail-line1-holiday-20260901.xlsx','holiday',1),
 ('korail-line4-20260822.xlsx',None,4),('korail-line75-20260822.xlsx',None,75),
]
data={}; unknown=set()
for file,day,fixed_line in sources:
 wb=openpyxl.load_workbook(ROOT/'data'/file,data_only=True)
 for sheet in wb:
  if fixed_line: line=fixed_line
  elif sheet.title.startswith('경의'): line=63
  elif sheet.title.startswith(('경춘','망우')): line=67
  else: continue  # ITX and Seohae do not supply express subway patterns here.
  period=day or ('weekday' if '평일' in sheet.title else 'holiday')
  valid={norm(n) for r in lines[str(line)]['routes'] for n in r['names']}
  line_data=data.setdefault(str(line),{'weekday':{},'holiday':{}})
  # Some line-1 sheets contain a second table to the right.
  headers=[c.column for c in sheet[4] if norm(c.value)=='열차번호']
  for hi,station_col in enumerate(headers):
   limit=headers[hi+1] if hi+1<len(headers) else sheet.max_column+1
   for col in range(station_col+1,limit):
    train=str(sheet.cell(4,col).value or '').strip()
    if not re.fullmatch(r'[KS]\d{3,5}',train): continue
    origin=norm(sheet.cell(2,col).value); terminal=norm(sheet.cell(3,col).value)
    stops=[]; passed=[]
    for row in range(5,sheet.max_row+1):
     label=sheet.cell(row,station_col).value
     if not label: continue
     station=norm(label)
     if station in {'마전','시흥연','시흥기'}: continue  # Non-passenger operating points.
     arrival=is_time(sheet.cell(row,col).value)
     departure=is_time(sheet.cell(row+1,col).value) if not sheet.cell(row+1,station_col).value else False
     if not arrival and not departure: continue
     if station not in valid:
      unknown.add((line,station));continue
     if arrival or (departure and station==origin):
      if station not in stops: stops.append(station)
     elif departure: passed.append(station)
    if not stops: continue
    key=re.sub(r'^[KS]0*','',train)
    record={'stops':stops,'passes':passed,'express':bool(passed),'origin':origin,'terminal':terminal}
    existing=line_data[period].get(key)
    if existing and existing!=record:
     # Duplicate tables must agree on each observed station, otherwise fail.
     assert not (set(existing['stops'])&set(passed) or set(existing['passes'])&set(stops)),(file,train)
     record['stops']=list(dict.fromkeys(existing['stops']+stops))
     record['passes']=list(dict.fromkeys(existing['passes']+passed))
     record['express']=existing['express'] or record['express']
    line_data[period][key]=record

if unknown:
 print('Unmapped station labels:',json.dumps(sorted(unknown),ensure_ascii=True))
 raise SystemExit(1)
# Only express records are needed. A missing train remains unknown, never a guessed stop.
for line,entry in data.items():
 for day in ['weekday','holiday']:
  entry[day]={k:v for k,v in entry[day].items() if v['express']}
 entry['stations']=sorted({s for day in ['weekday','holiday'] for v in entry[day].values() for s in v['stops']})
data['9']={'stations':['김포공항','마곡나루','가양','염창','당산','여의도','노량진','동작','고속터미널','신논현','선정릉','봉은사','종합운동장','석촌','올림픽공원','중앙보훈병원']}
(ROOT/'express-data.js').write_text('// Official Korail timetables (2026-09-01 / 2026-08-22) and Metro9. See data/README.md.\nconst EXPRESS_DATA = '+json.dumps(data,ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
print(json.dumps({k:{'weekday':len(v.get('weekday',{})),'holiday':len(v.get('holiday',{})),'stations':len(v['stations'])} for k,v in data.items()}))
